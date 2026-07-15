#!/usr/bin/env python3
"""Aggregate the HPC sweep results into the Reported-vs-Reproduced comparison.

Reads every results_*.csv produced by run_stepB.py, aggregates the 3 repetitions
per (language x model) into mean +/- SD, converts fractions -> percentages, and
writes:
  1. a tidy summary CSV (mean/SD per cell),
  2. the filled tracker spreadsheet (reproduced columns + auto deltas vs the paper),
  3. a printed comparison table.

Usage:
  python aggregate_results.py <folder-with-results-csvs> [tracker.xlsx]

Notes / guards:
  - run_stepB writes metrics as FRACTIONS (0-1); the paper reports PERCENTAGES. x100.
  - Rows with missing metrics (failed runs) are ignored, and reported as such.
  - Cells with fewer than 3 completed runs are flagged (the paper uses 3).
"""
import csv, glob, os, sys
import numpy as np

# exp_name looks like: "<lang>_<model>_r01"  (e.g. de_resnet18_r02)
LANGS = {"ar": "Arabic", "en": "English", "de": "German", "zh": "Mandarin",
         "ro": "Romanian", "ru": "Russian", "es": "Spanish"}
MODELS = {"resnet18": "ResNet-18", "resnet50": "ResNet-50", "septr": "SepTr",
          "ast": "AST", "wav2vec2": "wav2vec2.0", "whisper": "Whisper+MLP"}
METRICS = ["in_acc", "in_auc", "in_eer", "cross_acc", "cross_auc", "cross_eer"]


def parse_exp(name):
    parts = name.split("_")
    if len(parts) < 3:
        return None, None
    lang = parts[0]
    model = "_".join(parts[1:-1])          # model may contain no underscore, but be safe
    return (lang if lang in LANGS else None,
            model if model in MODELS else None)


def load(folder):
    cells = {}   # (lang, model) -> {metric: [values]}
    n_files = 0
    for path in sorted(glob.glob(os.path.join(folder, "results*.csv"))):
        n_files += 1
        with open(path) as f:
            for row in csv.DictReader(f):
                lang, model = parse_exp(row.get("exp_name", ""))
                if not lang or not model:
                    continue
                key = (lang, model)
                d = cells.setdefault(key, {m: [] for m in METRICS})
                for m in METRICS:
                    v = row.get(m, "")
                    if v not in ("", None):
                        try:
                            d[m].append(float(v))
                        except ValueError:
                            pass
    return cells, n_files


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else "."
    tracker = sys.argv[2] if len(sys.argv) > 2 else None

    cells, n_files = load(folder)
    if not cells:
        print(f"No usable rows found in {folder}/results*.csv "
              f"({n_files} file(s) scanned). Are the runs still going?")
        return

    # ---- tidy summary ----
    out = os.path.join(folder, "summary_mean_sd.csv")
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["language", "model", "n_runs"] +
                   [f"{m}_{s}" for m in METRICS for s in ("mean_pct", "sd_pct")])
        for (lang, model) in sorted(cells):
            d = cells[(lang, model)]
            n = max(len(d[m]) for m in METRICS)
            row = [LANGS[lang], MODELS[model], n]
            for m in METRICS:
                a = np.array(d[m]) * 100
                if len(a):
                    row += [round(a.mean(), 2), round(a.std(ddof=1), 2) if len(a) > 1 else 0.0]
                else:
                    row += ["", ""]
            w.writerow(row)
    print(f"Wrote {out}")

    # ---- printed comparison ----
    print(f"\n{'Language':10s} {'Model':12s} {'n':>2s}  "
          f"{'in-ACC':>14s} {'cross-ACC':>14s} {'cross-EER':>14s}")
    print("-" * 74)
    incomplete = []
    for (lang, model) in sorted(cells):
        d = cells[(lang, model)]
        n = max(len(d[m]) for m in METRICS)
        if n < 3:
            incomplete.append((LANGS[lang], MODELS[model], n))
        def ms(m):
            a = np.array(d[m]) * 100
            if not len(a):
                return "     -        "
            sd = a.std(ddof=1) if len(a) > 1 else 0.0
            return f"{a.mean():7.2f}+/-{sd:5.2f}"
        print(f"{LANGS[lang]:10s} {MODELS[model]:12s} {n:2d}  "
              f"{ms('in_acc')} {ms('cross_acc')} {ms('cross_eer')}")

    if incomplete:
        print("\nWARNING - cells with fewer than 3 completed runs (paper uses 3):")
        for lang, model, n in incomplete:
            print(f"  {lang} / {model}: only {n} run(s)")

    # ---- fill the tracker spreadsheet ----
    if tracker and os.path.exists(tracker):
        try:
            from openpyxl import load_workbook
            from openpyxl.comments import Comment
        except ImportError:
            print("\n(openpyxl not installed - skipping tracker fill)")
            return
        wb = load_workbook(tracker)
        ws = wb["Tracker"]
        # repro columns: in ACC=4, AUC=7, EER=10 ; cross ACC=13, AUC=16, EER=19
        COLS = {"in_acc": 4, "in_auc": 7, "in_eer": 10,
                "cross_acc": 13, "cross_auc": 16, "cross_eer": 19}
        filled = 0
        for r in range(3, ws.max_row + 1):
            lang_name = ws.cell(r, 1).value
            model_name = ws.cell(r, 2).value
            key = next((k for k in cells
                        if LANGS[k[0]] == lang_name and MODELS[k[1]] == model_name), None)
            if not key:
                continue
            d = cells[key]
            n = max(len(d[m]) for m in METRICS)
            note = [f"Reproduced on MeluXina (HPC), {n} run(s), paper batch sizes."]
            for m, col in COLS.items():
                a = np.array(d[m]) * 100
                if len(a):
                    ws.cell(r, col).value = round(float(a.mean()), 2)
                    sd = a.std(ddof=1) if len(a) > 1 else 0.0
                    note.append(f"{m}: {a.mean():.2f} +/- {sd:.2f}")
            ws.cell(r, 13).comment = Comment("\n".join(note), "repro")
            filled += 1
        wb.save(tracker)
        print(f"\nFilled {filled} row(s) in {tracker}")


if __name__ == "__main__":
    main()
